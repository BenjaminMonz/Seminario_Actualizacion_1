import csv
import os

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

from price_manager.services.audit_service import auditar_operacion


class ServicioReportePrecios:
  """Servicio para generar reportes Excel de comparación de precios."""

  @auditar_operacion("GENERACION_REPORTE_EXCEL")
  def generar_reporte_excel(
    self,
    ruta_alertas: str,
    ruta_reporte: str,
  ) -> str:
    """
    Genera un reporte Excel a partir del archivo de alertas.

    Args:
        ruta_alertas:
            Ruta del CSV de alertas generado previamente.
        ruta_reporte:
            Ruta donde se guardará el archivo Excel.

    Retorna:
        Ruta del archivo Excel generado.
    """
    if not os.path.exists(ruta_alertas):
      raise FileNotFoundError("No existe el archivo de alertas.")

    os.makedirs(
      os.path.dirname(ruta_reporte),
      exist_ok=True,
    )

    wb = Workbook()
    ws = wb.active
    ws.title = "Reporte precios"

    columnas = [
      "Producto",
      "Precio interno",
      "Precio web",
      "Diferencia",
      "Fecha de extracción",
    ]

    ws.append(columnas)

    with open(
      ruta_alertas,
      mode="r",
      encoding="utf-8",
    ) as archivo:
      reader = csv.DictReader(archivo)

      for fila in reader:
        ws.append([
          fila.get("producto"),
          float(fila.get("precio_interno")),
          float(fila.get("precio_web")),
          float(fila.get("diferencia")),
          fila.get("fecha_extraccion"),
        ])

    self._aplicar_formato(ws)

    wb.save(ruta_reporte)

    return ruta_reporte

  def _aplicar_formato(self, ws) -> None:
    """Aplica formato básico al reporte Excel."""
    encabezado_fill = PatternFill(
      start_color="D9EAF7",
      end_color="D9EAF7",
      fill_type="solid",
    )

    for celda in ws[1]:
      celda.font = Font(bold=True)
      celda.fill = encabezado_fill

    for columna in ws.columns:
      largo = max(
        len(str(celda.value))
        for celda in columna
        if celda.value is not None
      )
      letra = columna[0].column_letter
      ws.column_dimensions[letra].width = min(largo + 2, 40)

    for fila in ws.iter_rows(min_row=2):
      fila[1].number_format = "$ #,##0.00"
      fila[2].number_format = "$ #,##0.00"
      fila[3].number_format = "$ #,##0.00"
